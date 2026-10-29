"""Tạo file Google Sheet mẫu cho lịch chụp (đồng bộ 2 chiều với web nhân sự).

Chạy: python3 php-hostinger/tools/build-sheet-template.py <đường-dẫn.xlsx> [example.json]
Tiêu đề cột phải khớp agParse_ trong google-sheets/AnhGauSync.gs; lựa chọn Gói/Makeup khớp bookings.php.
Tab "Ví dụ" và "Mẫu trống" không có tên dạng "Tháng N" nên script không đồng bộ chúng.
"""
import datetime as dt
import json
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = 'Arial'
PACKAGES = ['399', '699', 'Couple 999', '1499', '2999', 'Team 349']
STAFF = ['HUY', 'CƯỜNG', 'TÍN']
MAKEUP = ['Nam', 'Nữ', 'Làm Tóc']
# (tiêu đề, độ rộng, định dạng số, căn giữa, ghi chú ở ô tiêu đề)
COLUMNS = [
    ('Ngày', 9, 'dd/mm', True, 'Ngày chụp dạng 05/11. Chỉ ghi ở dòng đầu của mỗi ngày; các dòng sau cùng ngày để trống (hoặc gộp ô).'),
    ('Giờ', 8, 'hh:mm', True, 'Giờ bắt đầu, dạng 10:00 hoặc 11:30.'),
    ('Tên khách', 30, '@', False, 'Tên khách, ghi luôn tiền cọc. VD: Chị Lan cọc 200k'),
    ('SĐT', 14, '@', True, 'Số điện thoại, giữ số 0 đầu.'),
    ('Gói', 12, '@', True, 'Để trống = 399/699, khách tới mới chọn (giữ 1 tiếng).\nCouple 999 / 1499 / 2999 = giữ 2 tiếng.\nTeam 349 = giữ theo "Giữ N tiếng" ở Ghi chú.'),
    ('Thợ chính', 11, '@', True, 'Tên đã ghép trên web (HUY, CƯỜNG…). Lịch thêm trên web tự ghi HUY.'),
    ('Thợ phụ', 11, '@', True, 'Lịch thêm trên web tự ghi CƯỜNG.'),
    ('Makeup', 14, '@', True, 'Nam / Nữ / Làm Tóc. Chọn nhiều được nếu bật "Cho phép nhiều lựa chọn" trong cài đặt dropdown.'),
    ('Tip', 10, '#,##0', True, 'Tiền tip khách cho (nếu có), VD 50000.'),
    ('Trạng thái', 11, '@', True, 'Khách hủy: chọn Hủy. Không xóa dòng (xóa dòng không xóa lịch trên web).'),
    ('Ghi chú', 34, '@', False, 'Ghi chú tự do. Team ghi "Giữ 3 tiếng" để khóa đủ giờ.'),
    ('_AG_BOOKING_ID', 38, '@', False, 'Cột ẩn: script tự ghi mã lịch. Không sửa, không xóa.'),
]
LAST = 'K'  # cột cuối người dùng thấy (Ghi chú); L = _AG_BOOKING_ID ẩn
line = Side(style='thin', color='D5DCE6')
BORDER = Border(left=line, right=line, top=line, bottom=line)
HEAD_FILL = PatternFill('solid', fgColor='FFFF00')


def data_sheet(wb, title, rows, blank_rows, explain=None, rule_rows=0):
    ws = wb.create_sheet(title)
    for c, (name, width, fmt, center, note) in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, name)
        cell.font = Font(name=FONT, bold=True, color='C00000', size=11)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
        cell.comment = Comment(note, 'Anh Gấu Studio', width=260, height=110)
        ws.column_dimensions[cell.column_letter].width = width
    ws.column_dimensions['L'].hidden = True
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    total = len(rows) + blank_rows
    for r in range(2, total + 2):
        for c, (_, _, fmt, center, _) in enumerate(COLUMNS, 1):
            cell = ws.cell(r, c)
            cell.font = Font(name=FONT, size=10)
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal='center' if center else 'left', vertical='center', wrap_text=c == 11)
            cell.border = BORDER
    for r, values in enumerate(rows, 2):
        for c, value in enumerate(values, 1):
            if value not in ('', None):
                ws.cell(r, c, value)
    end = max(total + 1, rule_rows)  # dropdown phủ rộng hơn vùng kẻ sẵn để dòng thêm sau vẫn có
    rules = [('E', PACKAGES, 'Chọn gói trong danh sách (để trống nếu khách chưa chọn 399/699).'),
             ('F', STAFF, 'Chọn tên thợ đã ghép trên web.'), ('G', STAFF, 'Chọn tên thợ đã ghép trên web.'),
             ('H', MAKEUP, 'Nam / Nữ / Làm Tóc.'), ('J', ['Hủy'], 'Chỉ chọn Hủy khi khách hủy.')]
    for col, options, message in rules:
        dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True,
                            showErrorMessage=True, errorStyle='warning', error=message, errorTitle='Ngoài danh sách')
        dv.add(f'{col}2:{col}{end}')
        ws.add_data_validation(dv)
    # Dòng đã hủy: chữ xám gạch ngang.
    ws.conditional_formatting.add(f'A2:{LAST}{end}', FormulaRule(formula=['$J2="Hủy"'], font=Font(strike=True, color='9AA3AE'), fill=PatternFill('solid', fgColor='EEF1F4')))
    if explain:
        col = len(COLUMNS) + 2  # chừa 1 cột trống sau cột ẩn
        head = ws.cell(1, col, 'Giải thích (chỉ có ở tab Ví dụ)')
        head.font = Font(name=FONT, bold=True, color='5F6B7A', size=10)
        head.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[head.column_letter].width = 62
        for r, text in enumerate(explain, 2):
            cell = ws.cell(r, col, text)
            cell.font = Font(name=FONT, italic=True, color='5F6B7A', size=10)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
    return ws


def build(path, example_json=None):
    wb = Workbook()
    guide = wb.active
    guide.title = 'Hướng dẫn'
    d1, d2, t = dt.date(2026, 11, 5), dt.date(2026, 11, 6), dt.time
    # Ví dụ: 05/11 và 06/11, không có dòng nào trùng giờ (script sẽ không tô đỏ).
    rows = [
        [d1, t(9, 0), 'Chị Lan cọc 200k', '0901234567', '', 'HUY', 'CƯỜNG', 'Nữ', None, '', ''],
        ['', t(10, 0), 'Anh Minh & chị Thảo cọc 500k', '0912345678', 'Couple 999', 'HUY', 'CƯỜNG', 'Nữ', None, '', 'Chụp áo dài'],
        ['', t(13, 0), 'Chị Hoa cọc 300k', '0987654321', '1499', 'HUY', 'CƯỜNG', 'Làm Tóc', 50000, '', ''],
        ['', t(15, 30), 'Chị Vy', '0934567890', '', 'HUY', 'CƯỜNG', '', None, 'Hủy', 'Khách dời sang tháng sau'],
        [d2, t(9, 0), 'Team công ty ABC cọc 1tr', '0909888777', 'Team 349', 'HUY', 'CƯỜNG', '', None, '', 'Giữ 3 tiếng'],
        ['', t(9, 0), 'Thành viên 2 · team ABC', '', 'Team 349', 'HUY', 'CƯỜNG', 'Nữ', None, '', ''],
        ['', t(9, 0), 'Thành viên 3 · team ABC', '', 'Team 349', 'HUY', 'CƯỜNG', '', None, '', ''],
        ['', t(13, 0), 'Anh Tuấn cọc 200k', '0977111222', '2999', 'HUY', 'CƯỜNG', 'Nam', None, '', ''],
    ]
    explain = [
        'Gói để trống = 399/699, khách tới mới chọn → giữ 1 tiếng (09:00–10:00). Ngày chỉ ghi ở dòng đầu của ngày.',
        'Couple 999 giữ 2 tiếng (10:00–12:00). Ô Ngày trống = cùng ngày với dòng trên.',
        'Gói 1499 giữ 2 tiếng (13:00–15:00). Cột Tip ghi số tiền khách tip.',
        'Khách hủy: chọn Hủy ở Trạng thái (dòng tự mờ đi). Không xóa dòng. Giờ lẻ 15:30 ghi được.',
        'Team: dòng đầu ghi "Giữ 3 tiếng" ở Ghi chú → khóa 09:00–12:00 với khách khác.',
        'Các thành viên team điền thêm cùng giờ, gói Team 349: không bị báo trùng nhau.',
        '',
        'Gói 2999 giữ 2 tiếng (13:00–15:00).',
    ]
    ex = data_sheet(wb, 'Ví dụ', rows, 20, explain)
    ex.merge_cells('A2:A5')
    ex.merge_cells('A6:A9')
    data_sheet(wb, 'Mẫu trống', [], 500, rule_rows=1000)

    guide.column_dimensions['A'].width = 4
    guide.column_dimensions['B'].width = 18
    guide.column_dimensions['C'].width = 12
    guide.column_dimensions['D'].width = 70
    guide.column_dimensions['E'].width = 40

    def put(r, c, text, **font):
        cell = guide.cell(r, c, text)
        cell.font = Font(name=FONT, size=font.pop('size', 10), **font)
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        return cell
    put(1, 2, 'Mẫu Google Sheet lịch chụp · đồng bộ 2 chiều với web nhân sự Anh Gấu Studio', bold=True, size=14)
    put(2, 2, 'Lịch thêm trên web (Lịch chụp → ＋ Thêm lịch) được script tự ghi vào tab "Tháng N" theo đúng các cột dưới đây. '
              'Lịch điền tay trên Sheet cũng được web đọc theo các cột này.', color='5F6B7A')
    guide.merge_cells('B1:E1')
    guide.merge_cells('B2:E2')
    guide.row_dimensions[2].height = 30
    r = 4
    put(r, 2, 'Cách dùng', bold=True, size=12)
    steps = [
        'Mở Google Sheet lịch hiện tại → Tệp → Nhập → Tải lên file này → chọn "Chèn trang tính mới". Script chỉ chạy được trong đúng Sheet đã cài kết nối.',
        'Chuột phải tab "Mẫu trống" → Tạo bản sao → đổi tên đúng dạng "Tháng 11" (hoặc "Tháng 12"…). Tab "Ví dụ" chỉ để xem, script không đồng bộ.',
        'Cột Makeup: Dữ liệu → Xác thực dữ liệu → chọn quy tắc của cột Makeup → Tùy chọn nâng cao → bật "Cho phép nhiều lựa chọn" (để chọn Nam + Nữ…). Có thể chọn kiểu hiển thị dạng "Khối" cho dễ nhìn.',
        'Điền lịch từ dòng 2 trở xuống, mỗi khách một dòng. Không sửa dòng tiêu đề (dòng 1) và cột ẩn _AG_BOOKING_ID.',
        'Menu Anh Gấu Studio → Đồng bộ ngay để kiểm tra. Từ tháng sau, script tự tạo tab tháng mới bằng cách nhân bản tab tháng liền trước khi có lịch thêm trên web.',
    ]
    for i, text in enumerate(steps, 1):
        r += 1
        put(r, 2, f'Bước {i}', bold=True)
        put(r, 4, text)
        guide.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        guide.row_dimensions[r].height = 32
    r += 2
    put(r, 2, 'Các cột', bold=True, size=12)
    r += 1
    for c, text in enumerate(['Cột', 'Bắt buộc', 'Cách điền', 'Lịch thêm trên web tự ghi'], 2):
        cell = put(r, c, text, bold=True, color='C00000')
        cell.fill = HEAD_FILL
        cell.border = BORDER
    table = [
        ('Ngày', 'Có', 'dd/mm (VD 05/11). Chỉ ghi ở dòng đầu của mỗi ngày; các dòng sau để trống hoặc gộp ô.', 'Chèn đúng ngày, đúng thứ tự giờ'),
        ('Giờ', 'Có', 'HH:mm (VD 10:00, 11:30).', 'Giờ đã chọn'),
        ('Tên khách', 'Có', 'Tên + tiền cọc (VD Chị Lan cọc 200k).', 'Tên khách đã nhập'),
        ('SĐT', '', 'Giữ số 0 đầu.', 'SĐT'),
        ('Gói', '', 'Để trống = 399/699 khách tới chọn (1 tiếng). Couple 999 / 1499 / 2999 = 2 tiếng. Team 349 = theo "Giữ N tiếng" ở Ghi chú.', 'Trống (399/699) hoặc gói đã chọn'),
        ('Thợ chính / Thợ phụ', '', 'Tên đã ghép trên web (HUY, CƯỜNG, TÍN…). Đổi thợ thì sửa ở đây, web tự cập nhật.', 'HUY / CƯỜNG'),
        ('Makeup', '', 'Nam / Nữ / Làm Tóc, chọn nhiều được.', 'Lựa chọn đã tick'),
        ('Tip', '', 'Số tiền tip (VD 50000).', '(không ghi)'),
        ('Trạng thái', '', 'Chọn Hủy khi khách hủy. Không xóa dòng.', 'Hủy, khi hủy trên web'),
        ('Ghi chú', '', 'Ghi chú tự do. Team: ghi "Giữ 3 tiếng" để khóa đủ giờ.', 'Ghi chú + "Giữ N tiếng" (Team)'),
        ('_AG_BOOKING_ID', '', 'Cột ẩn, script tự ghi mã lịch. Không sửa, không xóa.', 'Mã lịch'),
    ]
    for row in table:
        r += 1
        for c, text in enumerate(row, 2):
            put(r, c, text, bold=c == 2).border = BORDER
        guide.row_dimensions[r].height = 30
    r += 2
    put(r, 2, 'Quy tắc', bold=True, size=12)
    notes = [
        'Tên tab đúng dạng "Tháng 10", "Tháng 11", "Tháng 1"… Script tự đồng bộ các tab từ tháng trước đến 6 tháng tới.',
        'Mỗi khung giờ 1 khách. Lịch chồng giờ (kể cả điền tay) bị script tô đỏ ô Giờ và ghi chú dòng bị trùng.',
        'Team: chỉ cần 1 dòng có "Giữ N tiếng"; thêm các dòng thành viên cùng giờ với gói Team 349.',
        'Được thêm cột khác hoặc đổi thứ tự cột, miễn giữ đúng tên tiêu đề ở dòng 1.',
    ]
    for text in notes:
        r += 1
        put(r, 2, '•')
        put(r, 3, text)
        guide.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        guide.row_dimensions[r].height = 18
    guide.sheet_view.showGridLines = False
    wb.save(path)
    if example_json:
        # Giá trị hiển thị của tab Ví dụ, để kiểm thử bằng chính script Sheet.
        shown = [[c for c, _, _, _, _ in COLUMNS]]
        for values in rows:
            out = []
            for v in values + ['']:
                if isinstance(v, dt.date):
                    out.append(v.strftime('%d/%m'))
                elif isinstance(v, dt.time):
                    out.append(v.strftime('%H:%M'))
                elif v is None:
                    out.append('')
                else:
                    out.append(str(v))
            shown.append(out)
        with open(example_json, 'w', encoding='utf-8') as f:
            json.dump(shown, f, ensure_ascii=False)


if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
